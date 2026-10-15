"""Structured data import (GraphLab Import): CSV / Excel rows -> staged triples, no LLM.

A *mapping* says which columns make which entities and how they relate:

    {"entities": [
        {"key": "asset", "type": "Asset", "scope": true},                      # the selected asset
        {"key": "signal", "type": "Signal", "id_column": "Tag", "name_column": "Description",
         "properties": {"unit": "Unit"}},
        {"key": "limit", "type": "OperatingLimit", "id_column": "Tag", "id_suffix": "normal",
         "name_template": "{Description} normal range", "properties": {"low": "Low", "high": "High"},
         "skip_if_empty": ["Low", "High"]}],
     "relationships": [
        {"from": "component", "type": "MONITORED_BY", "to": "signal"},
        {"from": "asset", "type": "MONITORED_BY", "to": "signal", "unless": "component"},
        {"from": "signal", "type": "HAS_LIMIT", "to": "limit"}]}

Entity options used by the Limits template (P14b): `id_suffix_columns` (the id ends with
those columns' values, lower case; the first must not be empty), `existing_only` (a row whose
entity is not in the graph is refused, nothing new is made) and `replace` (on approval, a value
that differs from the graph's replaces it; the old value is kept and restored when the triple
is deleted). An OperatingLimit's `kind` is one of LIMIT_KINDS, and its unit must be its signal's.

Every row yields one triple per relationship whose two ends are present. Ids
are canonical (`bsk:<type>:<scope>/<key>[/<suffix>]`), matched to existing
entities only on exact id / key / name (never fuzzily). Triples are staged as
CandidateTriple (source="structured") and reviewed like extracted ones.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from django.conf import settings
from django.db import transaction
from django.utils import timezone

from . import registry
from .extraction import normalize
from .identity import EntityIndex, Match
from .units import normalize_unit
from .ids import clean_part, is_valid_id, make_id
from . import evidence
from .driver import GraphUnavailable, session
from .models import CandidateTriple, DataFile, Evidence
from .schema import Schema, SchemaError

KINDS = {".csv": "csv", ".tsv": "csv", ".txt": "csv", ".xlsx": "xlsx", ".xlsm": "xlsx"}
PROPERTY_RE = re.compile(r"^[a-z][a-z0-9_]*$")
RESERVED_PROPERTIES = {"id", "name", "source", "created_by_triple", "created_at", "updated_at", "data_file"}
SAMPLE_ROWS = 10
LIMIT_KINDS = ("normal", "specification", "alarm", "trip")
EVIDENCE_CHARS = 1000


class ImportError_(ValueError):
    """Invalid file, mapping or request."""


# --- files -----------------------------------------------------------------------

def file_dir(df: DataFile) -> Path:
    return Path(settings.GRAPH_DATA_BASE) / str(df.id)


def file_path(df: DataFile) -> Path:
    return file_dir(df) / df.filename


def _key_for(filename: str) -> str:
    base = re.sub(r"[^A-Za-z0-9]+", "-", Path(filename).stem).strip("-").upper()[:100] or "DATA"
    key, n = base, 1
    while DataFile.objects.filter(key=key).exists():
        n += 1
        key = f"{base}-{n}"
    return key


def add_data_file(uploaded) -> tuple[DataFile, bool]:
    """Store an uploaded CSV/XLSX (deduplicated by content) and read its columns."""
    name = Path(uploaded.name).name
    kind = KINDS.get(Path(name).suffix.lower())
    if kind is None:
        raise ImportError_(f"{name}: unsupported file type (use .csv or .xlsx)")
    if uploaded.size > settings.GRAPH_IMPORT_MAX_UPLOAD_BYTES:
        raise ImportError_(f"{name}: file is larger than {settings.GRAPH_IMPORT_MAX_UPLOAD_BYTES // (1024 * 1024)} MB")
    content = uploaded.read()
    sha = hashlib.sha256(content).hexdigest()
    existing = DataFile.objects.filter(sha256=sha).first()
    if existing:
        return existing, False

    df = DataFile(filename=name, key=_key_for(name), kind=kind, size=len(content), sha256=sha)
    file_dir(df).mkdir(parents=True, exist_ok=True)
    file_path(df).write_bytes(content)
    try:
        if kind == "xlsx":
            df.sheets = sheet_names(file_path(df))
            df.sheet = df.sheets[0] if df.sheets else ""
        refresh_columns(df)
    except Exception as exc:
        shutil.rmtree(file_dir(df), ignore_errors=True)
        raise ImportError_(f"{name}: could not read the file ({exc})") from exc
    df.save()
    return df, True


def delete_data_file(df: DataFile) -> dict:
    """Remove the file; its approved triples are taken back out of the graph first."""
    from . import triples

    committed = list(df.triples.filter(status=CandidateTriple.Status.APPROVED))
    if committed:
        triples.delete(committed)
    count = df.triples.count()
    shutil.rmtree(file_dir(df), ignore_errors=True)
    df.delete()
    return {"triples_deleted": count + len(committed)}


# --- reading tables --------------------------------------------------------------

def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def sheet_names(path: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _raw_rows(df: DataFile) -> list[list[str]]:
    path = file_path(df)
    if df.kind == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[df.sheet] if df.sheet else wb[wb.sheetnames[0]]
            return [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    data = path.read_bytes()
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def read_table(df: DataFile) -> tuple[list[str], list[tuple[int, dict[str, str]]]]:
    """(columns, [(spreadsheet row number, {column: value})]); fully empty rows are dropped."""
    raw = _raw_rows(df)
    header_index = max(1, df.header_row) - 1
    if header_index >= len(raw):
        raise ImportError_(f"header row {df.header_row} is beyond the end of the file ({len(raw)} rows)")
    columns, seen = [], {}
    for i, name in enumerate(raw[header_index]):
        name = name or f"Column {i + 1}"
        seen[name] = seen.get(name, 0) + 1
        columns.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    while columns and columns[-1].startswith("Column ") and all(
            len(r) <= len(columns) - 1 or not r[len(columns) - 1] for r in raw[header_index + 1:]):
        columns.pop()  # trailing empty columns
    rows = []
    for offset, values in enumerate(raw[header_index + 1:]):
        if not any(values):
            continue
        rows.append((header_index + 2 + offset, {c: (values[i] if i < len(values) else "") for i, c in enumerate(columns)}))
    return columns, rows


def refresh_columns(df: DataFile) -> None:
    columns, rows = read_table(df)
    if not columns:
        raise ImportError_("no columns found")
    df.columns, df.row_count = columns, len(rows)


def sample(df: DataFile, n: int = SAMPLE_ROWS) -> list[dict]:
    _, rows = read_table(df)
    return [{"row": number, "values": values} for number, values in rows[:n]]


# --- templates (column suggestions) -------------------------------------------------

def _find(columns: list[str], *synonyms: str, exclude: tuple[str, ...] = ()) -> str:
    """Best column for a list of synonyms: exact (normalized) first, then containment."""
    norm = {c: normalize(c) for c in columns if c not in exclude}
    for syn in synonyms:
        for column, n in norm.items():
            if n == syn:
                return column
    for syn in synonyms:
        for column, n in norm.items():
            if syn in n.split() or (len(syn) > 3 and syn in n):
                return column
    return ""


def _props(columns: list[str], spec: dict[str, tuple[str, ...]], exclude: tuple[str, ...] = ()) -> dict[str, str]:
    out = {}
    for prop, synonyms in spec.items():
        column = _find(columns, *synonyms, exclude=exclude)
        if column:
            out[prop] = column
    return out


COMPONENT_SYNONYMS = ("component", "equipment", "subsystem", "unit", "location", "area")
ASSET = {"key": "asset", "type": "Asset", "scope": True}


def _tag_list(columns: list[str]) -> dict:
    tag = _find(columns, "tag", "tag name", "tagname", "signal", "point", "variable", "name")
    name = _find(columns, "description", "desc", "signal name", "name", exclude=(tag,))
    component = _find(columns, "component", "equipment", "subsystem", "location", exclude=(tag, name))
    low = _find(columns, "low", "min", "lo", "lower", "low limit", "minimum")
    high = _find(columns, "high", "max", "hi", "upper", "high limit", "maximum")
    unit = _find(columns, "unit", "units", "uom", "eng unit", "engineering unit")
    # The PLC as a source (P14a) adds the OPC node of each signal and what a limit is based on.
    opc_node = _find(columns, "opc node", "node id", "nodeid", "opc node id")
    opc_slot = _find(columns, "opc slot", "slot")
    opc_type = _find(columns, "opc type")
    limit_kind = _find(columns, "limit kind")
    limit_basis = _find(columns, "limit basis", "basis")
    signal = {"key": "signal", "type": "Signal", "id_column": tag, "name_column": name,
              "properties": _props(columns, {"unit": ("unit", "units", "uom", "eng unit", "engineering unit"),
                                             "data_type": ("data type", "datatype", "type"),
                                             "description": ("description", "desc"),
                                             "historian_ref": ("historian", "historian ref", "source", "address")},
                                   exclude=(tag, component, opc_node, opc_slot, opc_type, limit_kind, limit_basis))}
    entities, relationships = [dict(ASSET)], []
    if component:
        entities.append({"key": "component", "type": "Component", "id_column": component, "name_column": component})
        relationships.append({"from": "component", "type": "MONITORED_BY", "to": "signal"})
    entities.append(signal)
    relationships.append({"from": "asset", "type": "MONITORED_BY", "to": "signal", **({"unless": "component"} if component else {})})
    if low or high:
        limit_props = {k: v for k, v in (("low", low), ("high", high), ("unit", unit), ("kind", limit_kind), ("basis", limit_basis)) if v}
        entities.append({"key": "limit", "type": "OperatingLimit", "id_column": tag, "id_suffix": "normal",
                         "name_template": f"{{{name or tag}}} normal range", "properties": limit_props,
                         "skip_if_empty": [c for c in (low, high) if c]})
        relationships.append({"from": "signal", "type": "HAS_LIMIT", "to": "limit"})
    if opc_node:
        entities.append({"key": "opc", "type": "OPCNode", "id_column": tag, "id_suffix": "opc", "name_column": opc_node,
                         "properties": {k: v for k, v in (("node_id", opc_node), ("slot", opc_slot), ("data_type", opc_type)) if v},
                         "skip_if_empty": [opc_node]})
        relationships.append({"from": "signal", "type": "IMPLEMENTED_AS", "to": "opc"})
    return {"entities": entities, "relationships": relationships}


def _alarm_list(columns: list[str]) -> dict:
    code = _find(columns, "code", "alarm code", "alarm", "alarm id", "tag", "id", "name")
    name = _find(columns, "name", "message", "text", "alarm text", "description", exclude=(code,))
    component = _find(columns, *COMPONENT_SYNONYMS, exclude=(code, name))
    alarm = {"key": "alarm", "type": "Alarm", "id_column": code, "name_column": name or code,
             "properties": {**({"code": code} if code else {}),
                            **_props(columns, {"severity": ("severity", "priority", "level", "class"),
                                               "description": ("description", "message", "text", "desc"),
                                               "plc_index": ("plc index", "fault index")},
                                     exclude=(code, component))}}
    entities = [dict(ASSET), alarm]
    relationships = [{"from": "asset", "type": "HAS_ALARM", "to": "alarm"}]
    if component:
        entities.append({"key": "component", "type": "Component", "id_column": component, "name_column": component})
        relationships.append({"from": "component", "type": "ASSOCIATED_WITH", "to": "alarm"})
    return {"entities": entities, "relationships": relationships}


def _bom(columns: list[str]) -> dict:
    parent = _find(columns, "parent", "parent item", "parent id", "parent component", "assembly", "belongs to")
    item = _find(columns, "item", "component", "part", "id", "item id", "part number", "name", exclude=(parent,))
    name = _find(columns, "description", "desc", "name", "item name", exclude=(parent, item))
    component = {"key": "component", "type": "Component", "id_column": item, "name_column": name or item,
                 "properties": _props(columns, {"component_type": ("type", "category", "class", "component type"),
                                                "part_number": ("part number", "part no", "pn", "partnumber"),
                                                "manufacturer": ("manufacturer", "vendor", "make", "supplier"),
                                                "quantity": ("qty", "quantity"),
                                                "description": ("description", "desc")},
                                      exclude=(parent, item))}
    entities, relationships = [dict(ASSET), component], []
    if parent:
        entities.append({"key": "parent", "type": "Component", "id_column": parent})
        relationships.append({"from": "parent", "type": "HAS_COMPONENT", "to": "component"})
    relationships.append({"from": "asset", "type": "HAS_COMPONENT", "to": "component", **({"unless": "parent"} if parent else {})})
    return {"entities": entities, "relationships": relationships}


def _limits(columns: list[str]) -> dict:
    """One row per signal and kind of limit (P14b). Only existing signals; a differing value is changed on approval."""
    tag = _find(columns, "tag", "tag name", "tagname", "signal", "historian column", "point", "variable")
    kind = _find(columns, "kind", "limit kind", "limit type", exclude=(tag,))
    name = _find(columns, "name", "limit name", exclude=(tag, kind))
    group = _find(columns, "group", "recipe", "product", exclude=(tag, kind, name))
    values = _props(columns, {"low": ("low", "lsl", "lower", "min"), "target": ("target", "nominal", "setpoint"),
                              "high": ("high", "usl", "upper", "max")}, exclude=(tag, kind, name, group))
    props = {**({"kind": kind} if kind else {}), **values, **({"group": group} if group else {}),
             **_props(columns, {"unit": ("unit", "units", "uom"), "basis": ("source", "basis", "origin")},
                      exclude=(tag, kind, name, group, *values.values()))}
    limit = {"key": "limit", "type": "OperatingLimit", "id_column": tag if kind else "",
             "id_suffix_columns": [c for c in (kind, group) if c], "replace": True,
             "properties": props, "skip_if_empty": list(values.values())}
    if name:
        limit["name_column"] = name
    elif tag and kind:
        limit["name_template"] = f"{{{tag}}} {{{kind}}} limits"
    return {"entities": [{"key": "signal", "type": "Signal", "id_column": tag, "existing_only": True}, limit],
            "relationships": [{"from": "signal", "type": "HAS_LIMIT", "to": "limit"}]}


TEMPLATES = (
    ("tag_list", "Tag list", "Signals monitored by the asset or its components, with optional operating limits.", _tag_list),
    ("limits", "Limits", "Limits of existing signals by kind (normal, specification, alarm, trip); a differing value is changed on approval.", _limits),
    ("alarm_list", "Alarm list", "Alarms of the asset, optionally associated with a component.", _alarm_list),
    ("bom", "Bill of materials", "Components of the asset, nested by a parent column.", _bom),
)


def _mapped_columns(mapping: dict) -> set[str]:
    used = set()
    for e in mapping["entities"]:
        used.update(c for c in (e.get("id_column"), e.get("name_column")) if c)
        used.update(c for c in (e.get("properties") or {}).values() if c)
        used.update(c for c in e.get("id_suffix_columns") or [] if c)
    return used


def suggestions(columns: list[str]) -> list[dict]:
    """Each built-in template filled in from the header names, best match first."""
    out = []
    for key, label, description, build_mapping in TEMPLATES:
        mapping = build_mapping(columns)
        out.append({"key": key, "label": label, "description": description, "mapping": mapping,
                    "matched_columns": len(_mapped_columns(mapping)),
                    "complete": all(e.get("scope") or e.get("id_column") for e in mapping["entities"])})
    order = {key: i for i, (key, *_) in enumerate(TEMPLATES)}
    return sorted(out, key=lambda s: (not s["complete"], -s["matched_columns"], order[s["key"]]))


# --- mapping validation ----------------------------------------------------------------

def validate_mapping(mapping, schema: Schema, columns: list[str], asset_id: str) -> list[str]:
    if not isinstance(mapping, dict):
        return ["mapping must be an object with entities and relationships"]
    entities, relationships = mapping.get("entities"), mapping.get("relationships")
    if not isinstance(entities, list) or not entities:
        return ["mapping needs at least one entity"]
    if not isinstance(relationships, list) or not relationships:
        return ["mapping needs at least one relationship"]
    errors, by_key, known = [], {}, set(columns)

    def column(e_key: str, what: str, name) -> None:
        if name and name not in known:
            errors.append(f"{e_key}: {what} column {name!r} is not in the file")

    for e in entities:
        key = str((e or {}).get("key") or "") if isinstance(e, dict) else ""
        if not key or key in by_key:
            errors.append(f"entity key {key!r} is missing or duplicated")
            continue
        by_key[key] = e
        if e.get("type") not in schema.entity_types:
            errors.append(f"{key}: unknown entity type {e.get('type')!r}")
        if e.get("scope"):
            if e.get("type") != "Asset":
                errors.append(f"{key}: only an Asset can be the selected asset")
            if not asset_id:
                errors.append(f"{key}: select an asset (the mapping uses the selected asset)")
            continue
        if not e.get("id_column"):
            errors.append(f"{key}: choose the key column")
        column(key, "key", e.get("id_column"))
        column(key, "name", e.get("name_column"))
        for placeholder in re.findall(r"\{([^{}]+)\}", e.get("name_template") or ""):
            column(key, "name template", placeholder)
        for c in e.get("skip_if_empty") or []:
            column(key, "skip-if-empty", c)
        for c in e.get("id_suffix_columns") or []:
            column(key, "id suffix", c)
        props = e.get("properties") or {}
        if not isinstance(props, dict):
            errors.append(f"{key}: properties must map property names to columns")
            continue
        for prop, col in props.items():
            if not PROPERTY_RE.match(str(prop)) or prop in RESERVED_PROPERTIES:
                errors.append(f"{key}: invalid property name {prop!r} (lower_snake_case, not reserved)")
            column(key, f"property {prop}", col)

    for r in relationships:
        r = r if isinstance(r, dict) else {}
        a, b = by_key.get(r.get("from")), by_key.get(r.get("to"))
        if a is None or b is None:
            errors.append(f"relationship {r.get('type')!r}: unknown entity {r.get('from')!r} or {r.get('to')!r}")
            continue
        if r.get("unless") and r["unless"] not in by_key:
            errors.append(f"relationship {r.get('type')!r}: unknown entity {r['unless']!r} in 'unless'")
        try:
            schema.check_relationship(a.get("type"), r.get("type"), b.get("type"))
        except SchemaError as exc:
            errors.append(str(exc))
    if asset_id and not is_valid_id(asset_id):
        errors.append(f"asset_id must be a canonical id (bsk:asset:...), got {asset_id!r}")
    return errors


# --- rows -> triples ---------------------------------------------------------------------

def _coerce(value: str):
    """Numbers become numbers (limits, quantities); everything else stays text."""
    try:
        return int(value)
    except ValueError:
        pass
    try:
        number = float(value)
    except ValueError:
        return value
    return number if math.isfinite(number) else value


@dataclass
class End:
    id: str
    type: str
    name: str
    existing: bool
    props: dict
    named: bool   # the name came from a name column/template (not just the key)
    match: str = ""
    candidates: list = field(default_factory=list)
    replace: bool = False   # a differing value is changed on approval (Limits template)


@dataclass
class BuildResult:
    triples: list[CandidateTriple] = field(default_factory=list)
    row_errors: list[dict] = field(default_factory=list)
    rows: int = 0
    skipped_rows: int = 0
    new_entities: int = 0
    existing_entities: int = 0
    possible_matches: int = 0   # entities that may be an existing one: resolved in review
    evidence: dict = field(default_factory=dict)   # id(triple) -> [Evidence] (one per row)

    def stats(self) -> dict:
        return {"rows": self.rows, "triples": len(self.triples), "skipped_rows": self.skipped_rows,
                "new_entities": self.new_entities, "existing_entities": self.existing_entities,
                "possible_matches": self.possible_matches,
                "row_errors": len(self.row_errors)}


def build(df: DataFile, mapping: dict, asset_id: str, index: EntityIndex | None = None) -> BuildResult:
    """Apply a (validated) mapping to the file's rows; returns unsaved CandidateTriples."""
    schema = registry.active_schema()
    index = index or EntityIndex.load()
    _, rows = read_table(df)
    if len(rows) > settings.GRAPH_IMPORT_MAX_ROWS:
        raise ImportError_(f"{len(rows)} rows: the limit is {settings.GRAPH_IMPORT_MAX_ROWS} rows per file")
    scope_key = asset_id.split(":", 2)[2] if asset_id else df.key
    scope_name = next((n.name for n in index.by_label.get("Asset", []) if n.id == asset_id), "") or scope_key
    result = BuildResult(rows=len(rows))
    staged: dict[tuple, CandidateTriple] = {}
    ends: dict[str, End] = {}   # one End per id across the file (names/properties are merged)
    current = _GraphProps()

    def refuse(number: int, message: str) -> None:
        result.row_errors.append({"row": number, "message": message})

    def entity(spec: dict, values: dict, number: int) -> End | None:
        if spec.get("scope"):
            return End(asset_id, "Asset", scope_name, index.has(asset_id), {}, True, "id" if index.has(asset_id) else "new")
        key = values.get(spec["id_column"], "")
        if not key or (spec.get("skip_if_empty") and not any(values.get(c) for c in spec["skip_if_empty"])):
            return None
        label = spec["type"]
        suffix_columns = spec.get("id_suffix_columns") or []
        if suffix_columns and not values.get(suffix_columns[0]):
            refuse(number, f"{spec['key']}: {suffix_columns[0]} is empty")
            return None
        try:
            parts = [clean_part(key)] + ([clean_part(spec["id_suffix"])] if spec.get("id_suffix") else [])
            parts += [clean_part(values[c].lower()) for c in suffix_columns if values.get(c)]
            proposed = make_id(label, *parts) if label == "Asset" else make_id(label, scope_key, *parts)
        except ValueError:
            result.row_errors.append({"row": number, "message": f"{spec['key']}: {key!r} can't be used as an id"})
            return None
        name = ""
        if spec.get("name_template"):
            name = re.sub(r"\{([^{}]+)\}", lambda m: values.get(m.group(1), ""), spec["name_template"]).strip()
        elif spec.get("name_column"):
            name = values.get(spec["name_column"], "")
        named = bool(name) and name != key
        if spec.get("id_suffix") or suffix_columns:   # derived entities (e.g. a signal's limit): exact id only
            match = Match(proposed, "id" if index.has(proposed) else "new")
        else:                       # the key column is an explicit tag; names / aliases, then possible matches
            match = index.resolve(label, name or key, scope_key, proposed=proposed, tag=key)
        if spec.get("existing_only") and not match.existing:
            refuse(number, f"{spec['key']}: {key!r} is not in the graph (this template does not add a {label})")
            return None
        node = index.get(match.id) if spec.get("existing_only") and not named else None
        if node:   # a key alone keeps the graph's name (no alias)
            name, named = node.name, True
        props = {p: _coerce(values[c]) for p, c in (spec.get("properties") or {}).items() if values.get(c)}
        if "unit" in props:
            props["unit"] = normalize_unit(props["unit"])
        if "subtype" in props and isinstance(props["subtype"], str):   # schema spelling when it is a known subtype
            props["subtype"] = schema.subtype_for(label, props["subtype"]) or props["subtype"]
        if label == "OperatingLimit" and "kind" in props:
            props["kind"] = str(props["kind"]).strip().lower()
            if props["kind"] not in LIMIT_KINDS:
                refuse(number, f"{spec['key']}: kind {props['kind']!r} is not one of {', '.join(LIMIT_KINDS)}")
                return None
        return End(match.id, label, (name or key)[:255], match.existing, props, named, match.kind, match.candidates,
                   replace=bool(spec.get("replace")))

    for number, values in rows:
        present: dict[str, End] = {}
        for spec in mapping["entities"]:
            end = entity(spec, values, number)
            if end is None:
                continue
            known = ends.get(end.id)
            if known is None:
                ends[end.id] = known = end
            else:  # same entity on another row: keep the first explicit name, add new properties
                if end.named and not known.named:
                    known.name, known.named = end.name, True
                for k, v in end.props.items():
                    known.props.setdefault(k, v)
            present[spec["key"]] = known
        emitted = 0
        for rel in mapping["relationships"]:
            a, b = present.get(rel["from"]), present.get(rel["to"])
            if a is None or b is None or (rel.get("unless") and rel["unless"] in present):
                continue
            if a.id == b.id:
                result.row_errors.append({"row": number, "message": f"{rel['type']}: {a.name!r} would point to itself"})
                continue
            if rel["type"] == "HAS_LIMIT" and b.props.get("unit") and a.existing:
                unit = current.get(a.id).get("unit")
                if unit and normalize_unit(str(unit)) != b.props["unit"]:
                    refuse(number, f"{b.name!r}: unit {b.props['unit']!r} is not the signal's ({unit!r})")
                    continue
            emitted += 1
            key = (a.id, rel["type"], b.id)
            excerpt = "; ".join(f"{c}: {v}" for c, v in values.items() if v)[:EVIDENCE_CHARS]
            if key not in staged:
                staged[key] = CandidateTriple(
                    data_file=df, source=CandidateTriple.Source.STRUCTURED, mode="schema",
                    subject_type=a.type, subject_id=a.id, predicate=rel["type"], object_type=b.type, object_id=b.id,
                    row_number=number, evidence_text=excerpt, schema_version=schema.version,
                )
                result.evidence[id(staged[key])] = []
            rows = result.evidence[id(staged[key])]
            rows.append(Evidence(source_kind=Evidence.Kind.STRUCTURED, data_file=df, row_number=number,
                                 extractor="column-mapping", schema_version=schema.version, excerpt=excerpt))
            staged[key].occurrences = len(rows)
        if not emitted:
            result.skipped_rows += 1
            if not any(e["row"] == number for e in result.row_errors):
                missing = [s["key"] for s in mapping["entities"] if s["key"] not in present]
                result.row_errors.append({"row": number, "message": f"no relationship built (missing: {', '.join(missing) or 'nothing'})"})

    used = set()
    for t in staged.values():  # final names/properties once every row has been read
        a, b = ends[t.subject_id], ends[t.object_id]
        t.subject_name, t.subject_existing, t.subject_props = a.name, a.existing, a.props
        t.object_name, t.object_existing, t.object_props = b.name, b.existing, b.props
        t.subject_match, t.subject_candidates = a.match, a.candidates
        t.object_match, t.object_candidates = b.match, b.candidates
        t.replaces = {e.id: c for e in (a, b) if (c := _changes(e, current))}
        used.update((a.id, b.id))
    result.triples = list(staged.values())
    result.existing_entities = sum(1 for i in used if ends[i].existing)
    result.possible_matches = sum(1 for i in used if ends[i].match == "possible")
    result.new_entities = len(used) - result.existing_entities - result.possible_matches
    return result


class _GraphProps(dict):
    """An existing entity's properties in the graph, read once per build (empty when the graph is off)."""

    def get(self, node_id: str) -> dict:
        if node_id not in self:
            try:
                with session() as s:
                    row = s.run("MATCH (n:Entity {id: $id}) RETURN properties(n) AS p", id=node_id).single()
            except GraphUnavailable:
                row = None
            self[node_id] = dict(row["p"]) if row else {}
        return self[node_id]


def _same(a, b) -> bool:
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return math.isclose(a, b, rel_tol=0, abs_tol=1e-9)
    return str(a) == str(b)


def _changes(end: End, current: _GraphProps) -> dict:
    """{property: [graph value, file value]} for an existing entity whose values the file changes on approval."""
    if not (end.replace and end.existing):
        return {}
    graph = current.get(end.id)
    return {k: [graph[k], v] for k, v in end.props.items() if k in graph and not _same(graph[k], v)}


def check(df: DataFile, mapping, asset_id: str) -> list[str]:
    columns, _ = read_table(df)
    return validate_mapping(mapping, registry.active_schema(), columns, asset_id or "")


@transaction.atomic
def stage(df: DataFile, mapping: dict, asset_id: str) -> dict:
    """Replace the file's pending triples with the mapping's result; reviewed triples are kept."""
    errors = check(df, mapping, asset_id)
    if errors:
        raise ImportError_("; ".join(errors))
    result = build(df, mapping, asset_id or "")
    reviewed = set(df.triples.exclude(status=CandidateTriple.Status.PENDING)
                   .values_list("subject_id", "predicate", "object_id"))
    fresh = [t for t in result.triples if (t.subject_id, t.predicate, t.object_id) not in reviewed]
    df.triples.filter(status=CandidateTriple.Status.PENDING).delete()
    evidence.attach(CandidateTriple.objects.bulk_create(fresh), result.evidence)
    evidence.drop_orphans(data_file=df)
    df.mapping, df.asset_id, df.staged_at = mapping, asset_id or "", timezone.now()
    df.save(update_fields=["mapping", "asset_id", "staged_at"])
    return {**result.stats(), "staged": len(fresh), "already_reviewed": len(result.triples) - len(fresh),
            "row_error_list": result.row_errors[:100]}
